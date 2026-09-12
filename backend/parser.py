"""Multi-format Document and Text parsing module for InternLoom Smart Shortlisting Engine.

Supports:
- PDF (.pdf) via PyMuPDF (fitz)
- DOCX (.docx) via python-docx
- TXT (.txt) with robust multi-encoding handling
- HTML (.html, .htm) via BeautifulSoup4 text extraction
"""

import io
import os
import re
import zipfile
from typing import Dict, List, Optional, Union
import fitz  # PyMuPDF
import docx  # python-docx
from bs4 import BeautifulSoup

from backend.cleaner import clean_text, extract_bullet_points
from backend.models import Resume, JobDescription


# Heading pattern definitions with flexible variants
SECTION_PATTERNS = {
    "skills": [
        r"(?:technical\s+)?skills",
        r"technical\s+expertise",
        r"core\s+competencies",
        r"technologies(?:\s+used)?",
        r"tools\s*(?:&|and)\s*technologies",
        r"key\s+skills",
        r"tech\s+stack",
        r"skills\s*&\s*abilities",
    ],
    "experience": [
        r"(?:work|professional|relevant|industry)\s+experience",
        r"experience",
        r"employment(?:\s+history)?",
        r"internships?(?:\s+experience)?",
        r"work\s+history",
        r"practical\s+experience",
    ],
    "projects": [
        r"(?:academic|personal|key|selected|technical)\s+projects",
        r"projects",
        r"project\s+work",
        r"portfolio",
        r"initiatives",
    ],
    "education": [
        r"education(?:al\s+background)?",
        r"academic(?:\s+background|\s+qualifications)?",
        r"qualifications",
        r"degrees?",
        r"educational\s+qualifications",
    ],
    "certifications": [
        r"certifications?",
        r"licenses\s*&\s*certifications",
        r"courses\s*&\s*certifications",
        r"accreditations?",
        r"achievements?(?:\s*&\s*awards)?",
        r"awards\s*&\s*achievements",
    ],
    "responsibilities": [
        r"responsibilities",
        r"key\s+responsibilities",
        r"duties",
        r"role\s+overview",
        r"what\s+you(?:'ll|\s+will)\s+do",
        r"what\s+you\s+will\s+be\s+doing",
        r"scope\s+of\s+work",
    ],
    "required_skills": [
        r"required\s+skills",
        r"must\s+haves?",
        r"minimum\s+qualifications",
        r"core\s+requirements",
        r"essential\s+skills",
        r"requirements",
        r"what\s+we\s+are\s+looking\s+for",
        r"basic\s+qualifications",
    ],
    "preferred_skills": [
        r"preferred\s+skills",
        r"nice\s+to\s+haves?",
        r"good\s+to\s+haves?",
        r"desired\s+skills",
        r"bonus\s+skills",
        r"preferred\s+qualifications",
        r"plus\s+points?",
    ],
    "education_requirements": [
        r"education\s+requirements?",
        r"eligibility\s+criteria",
        r"required\s+education",
    ],
    "experience_requirements": [
        r"experience\s+requirements?",
        r"required\s+experience",
    ],
}


# =========================================================================
# FORMAT-SPECIFIC EXTRACTORS
# =========================================================================

def parse_pdf(file_input: Union[str, bytes, io.BytesIO]) -> str:
    """Extracts raw text from a PDF file path or byte stream using PyMuPDF."""
    doc = None
    try:
        if isinstance(file_input, (bytes, bytearray)):
            doc = fitz.open(stream=file_input, filetype="pdf")
        elif isinstance(file_input, io.BytesIO):
            doc = fitz.open(stream=file_input.getvalue(), filetype="pdf")
        elif isinstance(file_input, str):
            doc = fitz.open(file_input)
        else:
            raise ValueError(f"Unsupported file_input type for PDF: {type(file_input)}")

        pages_text = []
        for page in doc:
            page_text = page.get_text()
            if page_text:
                pages_text.append(page_text)

        raw_text = "\n".join(pages_text)
        return clean_text(raw_text)
    except Exception as e:
        raise ValueError(f"PDF extraction failed: {str(e)}") from e
    finally:
        if doc:
            doc.close()


def parse_docx(file_input: Union[str, bytes, io.BytesIO]) -> str:
    """Extracts raw text from a DOCX file path or byte stream using python-docx."""
    try:
        if isinstance(file_input, (bytes, bytearray)):
            stream = io.BytesIO(file_input)
            doc = docx.Document(stream)
        elif isinstance(file_input, io.BytesIO):
            doc = docx.Document(file_input)
        elif isinstance(file_input, str):
            doc = docx.Document(file_input)
        else:
            raise ValueError(f"Unsupported file_input type for DOCX: {type(file_input)}")

        extracted_lines = []
        # Extract paragraph lines
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                extracted_lines.append(text)

        # Also extract table contents (often used in resumes)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    extracted_lines.append(" | ".join(row_cells))

        raw_text = "\n".join(extracted_lines)
        return clean_text(raw_text)
    except Exception as e:
        raise ValueError(f"DOCX extraction failed: {str(e)}") from e


def parse_txt(file_input: Union[str, bytes, io.BytesIO]) -> str:
    """Extracts text from a TXT file path or byte stream with multiple encoding fallbacks."""
    try:
        if isinstance(file_input, (bytes, bytearray)):
            byte_data = bytes(file_input)
        elif isinstance(file_input, io.BytesIO):
            byte_data = file_input.getvalue()
        elif isinstance(file_input, str):
            if os.path.exists(file_input):
                with open(file_input, "rb") as f:
                    byte_data = f.read()
            else:
                # Already a direct text string
                return clean_text(file_input)
        else:
            raise ValueError(f"Unsupported file_input type for TXT: {type(file_input)}")

        # Try multiple encodings
        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"):
            try:
                decoded = byte_data.decode(encoding)
                return clean_text(decoded)
            except (UnicodeDecodeError, LookupError):
                continue

        # Last resort: ignore errors
        return clean_text(byte_data.decode("utf-8", errors="ignore"))
    except Exception as e:
        raise ValueError(f"TXT extraction failed: {str(e)}") from e


def parse_html(file_input: Union[str, bytes, io.BytesIO]) -> str:
    """Extracts human-readable text from HTML markup, removing script and style tags."""
    try:
        if isinstance(file_input, (bytes, bytearray)):
            html_content = bytes(file_input).decode("utf-8", errors="ignore")
        elif isinstance(file_input, io.BytesIO):
            html_content = file_input.getvalue().decode("utf-8", errors="ignore")
        elif isinstance(file_input, str):
            if os.path.exists(file_input):
                with open(file_input, "r", encoding="utf-8", errors="ignore") as f:
                    html_content = f.read()
            else:
                html_content = file_input
        else:
            raise ValueError(f"Unsupported file_input type for HTML: {type(file_input)}")

        soup = BeautifulSoup(html_content, "html.parser")

        # Strip scripts, styles, metadata, and header tags
        for element in soup(["script", "style", "meta", "link", "noscript", "svg", "head"]):
            element.decompose()

        # Extract structured text with newlines between blocks
        text = soup.get_text(separator="\n")
        return clean_text(text)
    except Exception as e:
        raise ValueError(f"HTML extraction failed: {str(e)}") from e


# =========================================================================
# UNIVERSAL FORMAT DETECTOR & DISPATCHER
# =========================================================================

def detect_file_format(
    file_bytes: bytes,
    filename: Optional[str] = None
) -> str:
    """Detects file format based on magic bytes and filename extension.
    
    Returns one of: 'pdf', 'docx', 'txt', 'html'.
    Raises ValueError for unsupported formats.
    """
    ext = ""
    if filename:
        ext = os.path.splitext(filename)[1].lower()

    # 1. Magic byte checks
    if file_bytes.startswith(b"%PDF"):
        return "pdf"

    # DOCX is a zip archive containing word/document.xml or [Content_Types].xml
    if file_bytes.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                names = z.namelist()
                if any("word/" in n for n in names) or "[Content_Types].xml" in names:
                    return "docx"
        except zipfile.BadZipFile:
            pass

    # HTML check: look for html tags or doctype
    prefix_sample = file_bytes[:1000].lower()
    if (
        b"<!doctype html" in prefix_sample
        or b"<html" in prefix_sample
        or b"<body" in prefix_sample
        or b"<div" in prefix_sample
        or b"<h1" in prefix_sample
        or ext in [".html", ".htm"]
    ):
        # Only classify as html if it actually has markup tags
        if b"<" in prefix_sample and b">" in prefix_sample:
            return "html"

    # 2. Filename extension fallback
    if ext == ".pdf":
        return "pdf"
    elif ext == ".docx":
        return "docx"
    elif ext in [".html", ".htm"]:
        return "html"
    elif ext in [".txt", ".text", ".md", ".rtf", ""]:
        return "txt"

    # Default fallback to plain text if readable
    try:
        file_bytes[:512].decode("utf-8")
        return "txt"
    except UnicodeDecodeError:
        raise ValueError(
            f"Unsupported file type for '{filename or 'uploaded document'}'. "
            "Please upload PDF (.pdf), Word (.docx), Plain Text (.txt), or HTML (.html)."
        )


def extract_text_from_file(
    file_input: Union[str, bytes, io.BytesIO],
    filename: Optional[str] = None
) -> str:
    """Universal text extractor routing input through the correct format engine."""
    # Handle direct string inputs
    if isinstance(file_input, str):
        if not os.path.exists(file_input):
            # Direct raw text string
            if file_input.strip().startswith("<") and ("</html>" in file_input or "</div>" in file_input):
                return parse_html(file_input)
            return clean_text(file_input)
        
        # It's an existing file path
        filename = filename or os.path.basename(file_input)
        with open(file_input, "rb") as f:
            file_bytes = f.read()
    elif isinstance(file_input, io.BytesIO):
        file_bytes = file_input.getvalue()
    elif isinstance(file_input, (bytes, bytearray)):
        file_bytes = bytes(file_input)
    else:
        raise ValueError(f"Unsupported file input type: {type(file_input)}")

    if not file_bytes or not file_bytes.strip():
        raise ValueError(f"Uploaded file '{filename or 'document'}' is empty.")

    fmt = detect_file_format(file_bytes, filename=filename)

    if fmt == "pdf":
        text = parse_pdf(file_bytes)
    elif fmt == "docx":
        text = parse_docx(file_bytes)
    elif fmt == "html":
        text = parse_html(file_bytes)
    elif fmt == "txt":
        text = parse_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported format: {fmt}")

    if not text.strip():
        raise ValueError(
            f"Could not extract any readable text from '{filename or 'document'}'. "
            "Please ensure the file contains valid text."
        )

    return text


# =========================================================================
# SECTIONING AND STRUCTURE EXTRACTION
# =========================================================================

def detect_sections(text: str) -> Dict[str, str]:
    """Segment resume or JD text into recognized section buckets."""
    lines = text.split("\n")
    sections: Dict[str, List[str]] = {}
    current_section = "header"
    sections[current_section] = []

    # Build regex mapping for all defined section types
    section_regexes = {}
    for sec_type, patterns in SECTION_PATTERNS.items():
        pattern_str = r"^(?:[0-9]+[.\-)]\s*)?(?:[-*#•\s]*)\b(?:" + "|".join(patterns) + r")\b\s*[:\-—]*\s*$"
        section_regexes[sec_type] = re.compile(pattern_str, re.IGNORECASE)

    for line in lines:
        stripped = line.strip()
        if not stripped:
            if current_section in sections:
                sections[current_section].append("")
            continue

        matched_section_type = None
        for sec_type, regex in section_regexes.items():
            if regex.match(stripped):
                matched_section_type = sec_type
                break

        if matched_section_type:
            current_section = matched_section_type
            if current_section not in sections:
                sections[current_section] = []
        else:
            sections[current_section].append(stripped)

    # Convert list of lines back to string for each section
    return {k: "\n".join(v).strip() for k, v in sections.items() if "\n".join(v).strip()}


def extract_skills_from_text(skills_text: str) -> List[str]:
    """Splits skill section text into individual skill tokens."""
    if not skills_text:
        return []

    # Delimiters commonly used in skill sections
    raw_tokens = re.split(r"[,;|•\n/·]|\s+-\s+", skills_text)
    skills = []
    for token in raw_tokens:
        clean_token = re.sub(r"^[-*•\s]+", "", token).strip()
        # Filter out numbers, empty, or multi-sentence lines
        if clean_token and len(clean_token) <= 40 and not clean_token.isdigit():
            # Avoid general stop words
            if not re.match(r"^(and|or|with|including|such as)$", clean_token, re.IGNORECASE):
                skills.append(clean_token)
    return skills


def parse_resume(
    input_data: Union[str, bytes, io.BytesIO],
    candidate_id: str = "C01",
    filename: Optional[str] = None
) -> Resume:
    """Parses a resume from any supported document format (PDF, DOCX, TXT, HTML) or raw text."""
    raw_text = extract_text_from_file(input_data, filename=filename)
    sections = detect_sections(raw_text)

    # Candidate Name heuristic: Look at first non-empty line of header
    name = f"Candidate {candidate_id}"
    header_text = sections.get("header", "")
    if header_text:
        first_line = header_text.split("\n")[0].strip()
        # Clean title artifacts from header lines (e.g. "Aarav Sharma - Resume")
        first_line = re.sub(r"(?i)\s*[-—|•/]\s*(?:resume|curriculum vitae|cv|profile)\s*$", "", first_line).strip()
        first_line = re.sub(r"[^a-zA-Z0-9\s.,]", " ", first_line).strip()
        first_line = re.sub(r"\s+", " ", first_line)
        if first_line and not re.search(r"[@\d+]", first_line) and len(first_line) < 50:
            name = first_line
    elif filename:
        # Fallback to sanitized filename
        base_name = re.sub(r"\.(pdf|docx|txt|html|htm|text|rtf)$", "", filename, flags=re.IGNORECASE)
        base_name = re.sub(r"(?i)^resume[-_]?", "", base_name)
        base_name = re.sub(r"[-_]", " ", base_name).strip().title()
        if base_name:
            name = base_name

    # Extract skills
    skills_text = sections.get("skills", "")
    skills = extract_skills_from_text(skills_text)

    # Extract experience, projects, education, certifications
    experience = extract_bullet_points(sections.get("experience", ""))
    projects = extract_bullet_points(sections.get("projects", ""))
    education = extract_bullet_points(sections.get("education", ""))
    certifications = extract_bullet_points(sections.get("certifications", ""))

    return Resume(
        candidate_id=candidate_id,
        name=name,
        skills=skills,
        experience=experience,
        projects=projects,
        education=education,
        certifications=certifications,
        raw_text=raw_text
    )


def parse_job_description(
    input_data: Union[str, bytes, io.BytesIO],
    default_title: str = "Junior Full Stack Developer Intern",
    filename: Optional[str] = None
) -> JobDescription:
    """Parses a Job Description from any supported document format (PDF, DOCX, TXT, HTML) or raw text."""
    raw_text = extract_text_from_file(input_data, filename=filename)
    sections = detect_sections(raw_text)

    # Extract job title
    job_title = default_title
    header_text = sections.get("header", "")
    if header_text:
        lines = [line.strip() for line in header_text.split("\n") if line.strip()]
        if lines:
            for candidate_line in lines[:3]:
                if any(w in candidate_line.lower() for w in ["developer", "engineer", "intern", "role", "full stack", "analyst"]):
                    job_title = candidate_line
                    break

    # Extract required and preferred skills
    req_skills_raw = sections.get("required_skills", "")
    if not req_skills_raw and "skills" in sections:
        # Fallback if JD just lists "Skills"
        req_skills_raw = sections.get("skills", "")

    required_skills = extract_skills_from_text(req_skills_raw)
    if not required_skills:
        # Try extracting bullet items
        required_skills = extract_bullet_points(req_skills_raw)

    pref_skills_raw = sections.get("preferred_skills", "")
    preferred_skills = extract_skills_from_text(pref_skills_raw)
    if not preferred_skills:
        preferred_skills = extract_bullet_points(pref_skills_raw)

    responsibilities = extract_bullet_points(sections.get("responsibilities", ""))
    education_req = extract_bullet_points(sections.get("education_requirements", "") or sections.get("education", ""))
    experience_req = extract_bullet_points(sections.get("experience_requirements", ""))

    description_parts = []
    if "header" in sections:
        description_parts.append(sections["header"])
    if "responsibilities" in sections:
        description_parts.append(sections["responsibilities"])

    description = "\n\n".join(description_parts)

    return JobDescription(
        job_title=job_title,
        company="TechNova Solutions",
        required_skills=required_skills,
        preferred_skills=preferred_skills,
        responsibilities=responsibilities,
        education_requirements=education_req,
        experience_requirements=experience_req,
        description=description,
        raw_text=raw_text
    )

